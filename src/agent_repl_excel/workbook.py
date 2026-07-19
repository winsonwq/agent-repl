from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from agent_repl_sdk import ctx


ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}


class ExcelWorkbook:
    def __init__(self, path: Path):
        self.path = path.resolve()
        self.book = load_workbook(self.path)
        self._dirty = False
        ctx.events.emit("workbook_opened", path=str(self.path), sheets=self.book.sheetnames)

    @classmethod
    def open_input(cls, name: str, working_name: str | None = None) -> "ExcelWorkbook":
        source_path = ctx.inputs.path(name)
        if not source_path.is_file():
            workspace = ctx.paths["inputs"].resolve().parent
            candidate = Path(name)
            candidate = candidate.resolve() if candidate.is_absolute() else (workspace / candidate).resolve()
            if candidate != workspace and workspace not in candidate.parents:
                raise ValueError(f"Workbook must be inside the current workspace: {name}")
            if not candidate.is_file():
                raise FileNotFoundError(f"No such workbook in the workspace or inputs directory: {name}")
            source = candidate
        else:
            source = source_path
        working = ctx.working.copy_from(source, working_name or Path(name).name)
        return cls(working.path)

    open = open_input

    @classmethod
    def open_working(cls, name: str) -> "ExcelWorkbook":
        return cls(ctx.working.get(name).path)

    def list_sheets(self) -> list[str]:
        return list(self.book.sheetnames)

    def _sheet_and_address(self, qualified: str):
        if "!" not in qualified:
            raise ValueError("Use a sheet-qualified address such as Sales!A1")
        sheet_name, address = qualified.rsplit("!", 1)
        sheet_name = sheet_name.strip("'")
        return self.book[sheet_name], address

    def read_cell(self, qualified: str) -> Any:
        sheet, address = self._sheet_and_address(qualified)
        value = sheet[address].value
        ctx.events.emit("spreadsheet_cell_read", workbook=str(self.path), address=qualified)
        return value

    def read_range(self, qualified: str) -> list[list[Any]]:
        sheet, address = self._sheet_and_address(qualified)
        min_col, min_row, max_col, max_row = range_boundaries(address)
        values = [
            [sheet.cell(row=row, column=col).value for col in range(min_col, max_col + 1)]
            for row in range(min_row, max_row + 1)
        ]
        ctx.events.emit("spreadsheet_range_read", workbook=str(self.path), address=qualified)
        return values

    def set_value(self, qualified: str, value: Any) -> None:
        sheet, address = self._sheet_and_address(qualified)
        sheet[address] = value
        self._dirty = True
        ctx.events.emit("spreadsheet_cell_write", workbook=str(self.path), address=qualified, kind="value")

    def set_formula(self, qualified: str, formula: str) -> None:
        if not formula.startswith("="):
            raise ValueError("Formula must begin with '='")
        sheet, address = self._sheet_and_address(qualified)
        sheet[address] = formula
        self._dirty = True
        ctx.events.emit("spreadsheet_cell_write", workbook=str(self.path), address=qualified, kind="formula")

    def save_working(self) -> Path:
        calculation = getattr(self.book, "calculation", None)
        if calculation is not None:
            calculation.fullCalcOnLoad = True
            calculation.forceFullCalc = True
        self.book.save(self.path)
        self._dirty = False
        ctx.events.emit("workbook_saved", path=str(self.path))
        return self.path

    def recalculate(self, soffice: str | None = None) -> dict[str, Any]:
        self.save_working()
        executable = soffice or os.environ.get("AGENT_REPL_SOFFICE") or shutil.which("soffice") or shutil.which("libreoffice")
        if not executable:
            return {"recalculated": False, "engine": None, "reason": "LibreOffice not available; workbook marked for full calculation on open"}
        with tempfile.TemporaryDirectory(prefix="agent-repl-calc-") as temporary:
            temporary_path = Path(temporary)
            source_dir = temporary_path / "source"
            output_dir = temporary_path / "output"
            source_dir.mkdir()
            output_dir.mkdir()
            staged = source_dir / self.path.name
            shutil.copy2(self.path, staged)
            completed = subprocess.run(
                [executable, "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(staged)],
                capture_output=True,
                text=True,
                timeout=60,
            )
            calculated = output_dir / self.path.name
            if completed.returncode != 0 or not calculated.exists():
                raise RuntimeError(f"LibreOffice recalculation failed: {completed.stderr or completed.stdout}")
            shutil.copy2(calculated, self.path)
        self.book = load_workbook(self.path)
        self._dirty = False
        ctx.events.emit("workbook_recalculated", path=str(self.path), engine="libreoffice")
        return {"recalculated": True, "engine": "libreoffice"}

    def cached_value(self, qualified: str) -> Any:
        data_book = load_workbook(self.path, data_only=True)
        sheet_name, address = qualified.rsplit("!", 1)
        return data_book[sheet_name.strip("'")][address].value

    def scan_formula_errors(self) -> list[dict[str, Any]]:
        errors: list[dict[str, Any]] = []
        # Inspect cached calculation results as well as literal error values.
        # Formula cells themselves contain strings such as "=A1/B1" in the
        # normal workbook view, so only the data-only view reveals #DIV/0!, etc.
        data_book = load_workbook(self.path, data_only=True)
        for sheet in data_book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in ERROR_VALUES:
                        errors.append({"sheet": sheet.title, "cell": cell.coordinate, "error": cell.value})
        ctx.events.emit("workbook_validated", path=str(self.path), formula_errors=len(errors))
        return errors

    def save_output(self, name: str, title: str | None = None) -> dict[str, Any]:
        # Saving a LibreOffice-recalculated workbook again with openpyxl strips
        # cached formula results. Save only when edits are still pending.
        if self._dirty:
            self.save_working()
        output_path = ctx.paths["outputs"] / name
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.path, output_path)
        return ctx.outputs.publish_file(
            output_path,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            title=title or name,
        )


class ExcelRuntime:
    def open_input(self, name: str, working_name: str | None = None) -> ExcelWorkbook:
        return ExcelWorkbook.open_input(name, working_name)

    open = open_input

    def capabilities(self) -> dict[str, Any]:
        return {
            "openpyxl": True,
            "libreoffice": bool(os.environ.get("AGENT_REPL_SOFFICE") or shutil.which("soffice") or shutil.which("libreoffice")),
        }
