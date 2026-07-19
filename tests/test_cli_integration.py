from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import stat
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FIXTURE = PROJECT_ROOT / "tests" / "fixtures" / "sales_model.xlsx"


class CliIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temporary.name)
        cls.state_dir = cls.root / "state"
        cls.workspace = cls.root / "workspace"
        (cls.workspace / "inputs").mkdir(parents=True)
        shutil.copy2(FIXTURE, cls.workspace / "inputs" / FIXTURE.name)
        shutil.copy2(FIXTURE, cls.workspace / "report-at-root.xlsx")
        cls.environment = os.environ.copy()
        cls.environment.update(
            {
                "AGENT_TASK_ID": "e2e-test",
                "AGENT_REPL_STATE_DIR": str(cls.state_dir),
                "AGENT_REPL_WORKDIR": str(cls.workspace),
                "AGENT_REPL_OUTPUT": "json",
                "AGENT_REPL_TEST_SECRET": "must-not-enter-kernel",
            }
        )

    @classmethod
    def tearDownClass(cls):
        cls.run_cli("clean", check=False)
        cls.temporary.cleanup()

    @classmethod
    def run_cli(cls, *arguments: str, check: bool = True):
        completed = subprocess.run(
            [sys.executable, "-m", "agent_repl.cli", *arguments],
            cwd=PROJECT_ROOT,
            env=cls.environment,
            capture_output=True,
            text=True,
            timeout=90,
        )
        if check and completed.returncode != 0:
            raise AssertionError(f"CLI failed ({completed.returncode})\nstdout={completed.stdout}\nstderr={completed.stderr}")
        payload = json.loads(completed.stdout) if completed.stdout.strip() else None
        return completed, payload

    def test_01_python_state_persists_across_cli_calls(self):
        _, first = self.run_cli("py@state", "x = 41")
        _, second = self.run_cli("py@state", "x + 1")
        self.assertTrue(first["data"]["created"])
        self.assertTrue(second["data"]["state_preserved"])
        result_events = [event for event in second["data"]["events"] if event["type"] == "result"]
        self.assertEqual(result_events[-1]["text"], "42")

    def test_02_dataframe_state_persists(self):
        self.run_cli("data@analysis", "df = pd.DataFrame({'region':['East','West'], 'revenue':[1200,1500]})")
        _, payload = self.run_cli("data@analysis", "int(df['revenue'].sum())")
        results = [event for event in payload["data"]["events"] if event["type"] == "result"]
        self.assertEqual(results[-1]["text"], "2700")

    def test_03_excel_end_to_end(self):
        code = """
wb = excel.open_input('sales_model.xlsx')
before = wb.read_cell('Sales!B6')
wb.set_value('Sales!B6', 1500)
wb.set_formula('Sales!D6', '=B6-C6')
calc = wb.recalculate()
errors = wb.scan_formula_errors()
published = wb.save_output('sales_model_updated.xlsx', title='Updated sales model')
{'before': before, 'after': wb.read_cell('Sales!B6'), 'cached_profit': wb.cached_value('Sales!D6'), 'calc': calc, 'errors': errors, 'published': published}
"""
        _, payload = self.run_cli("excel@finance", code)
        result_events = [event for event in payload["data"]["events"] if event["type"] == "result"]
        self.assertTrue(result_events)

        output = self.workspace / "outputs" / "sales_model_updated.xlsx"
        self.assertTrue(output.exists())
        workbook = load_workbook(output, data_only=False)
        self.assertEqual(workbook["Sales"]["B6"].value, 1500)
        self.assertEqual(workbook["Sales"]["D6"].value, "=B6-C6")
        calculated = load_workbook(output, data_only=True)
        self.assertEqual(calculated["Sales"]["D6"].value, 720)

        # A second CLI process reuses the in-memory wb object.
        _, reuse = self.run_cli("excel@finance", "wb.read_cell('Sales!B6')")
        reuse_results = [event for event in reuse["data"]["events"] if event["type"] == "result"]
        self.assertEqual(reuse_results[-1]["text"], "1500")

    def test_04_error_is_structured(self):
        completed, payload = self.run_cli("py@errors", "raise ValueError('expected failure')", check=False)
        self.assertEqual(completed.returncode, 1)
        self.assertEqual(payload["status"], "error")
        self.assertEqual(payload["error"]["code"], "CODE_EXECUTION_ERROR")
        self.assertFalse(payload["data"]["ok"])
        errors = [event for event in payload["data"]["events"] if event["type"] == "error"]
        self.assertEqual(errors[-1]["name"], "ValueError")

    def test_05_kernel_environment_is_sanitized(self):
        _, payload = self.run_cli("py@environment", "import os; 'AGENT_REPL_TEST_SECRET' in os.environ")
        results = [event for event in payload["data"]["events"] if event["type"] == "result"]
        self.assertEqual(results[-1]["text"], "False")

    def test_06_timeout_interrupts_and_preserves_responsive_kernel(self):
        completed, payload = self.run_cli(
            "py@timeout",
            "import time; time.sleep(10)",
            "--timeout",
            "0.3",
            check=False,
        )
        self.assertEqual(completed.returncode, 4)
        self.assertEqual(payload["error"]["code"], "EXECUTION_TIMEOUT")
        self.assertTrue(payload["data"]["session_recovered"])
        _, reuse = self.run_cli("py@timeout", "21 * 2")
        results = [event for event in reuse["data"]["events"] if event["type"] == "result"]
        self.assertEqual(results[-1]["text"], "42")

    def test_07_excel_can_open_source_from_workspace_root(self):
        code = "wb_root = excel.open('report-at-root.xlsx'); wb_root.read_cell('Sales!B4')"
        _, payload = self.run_cli("excel@root-file", code)
        results = [event for event in payload["data"]["events"] if event["type"] == "result"]
        self.assertEqual(results[-1]["text"], "1000")
        self.assertTrue((self.workspace / "working" / "report-at-root.xlsx").is_file())

    def test_08_kernel_control_files_are_owner_only(self):
        _, payload = self.run_cli("sessions")
        self.assertTrue(payload["data"])
        record = payload["data"][0]
        self.assertEqual(stat.S_IMODE(Path(record["connection_file"]).stat().st_mode), 0o600)
        self.assertEqual(stat.S_IMODE(Path(record["context_file"]).stat().st_mode), 0o600)


if __name__ == "__main__":
    unittest.main()
